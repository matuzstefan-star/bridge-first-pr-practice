import datetime
import tempfile
import unittest
from pathlib import Path

from docx import Document

from bridge_analyzer.__main__ import archive, collect, main
from bridge_analyzer.analysis import analyze_board
from bridge_analyzer.docx_report import build_report, result_text
from bridge_analyzer.lin_tools import parse_lin

from .test_analysis import _deal, _lin, _simulate


def _lin_for(seed, strain="NT"):
    deal = _deal(seed)
    plays, _ = _simulate(deal, "S", strain, seed)
    return _lin(deal, plays, f"3{strain}")


class DailyTests(unittest.TestCase):
    def test_duplicates_are_one_hand_and_archive_moves_all_copies(self):
        with tempfile.TemporaryDirectory() as tmp:
            src, dst = Path(tmp, "new"), Path(tmp, "done")
            src.mkdir()
            (src / "1.lin").write_text(_lin_for(1), encoding="utf-8")
            (src / "1 (1).lin").write_text(_lin_for(1), encoding="utf-8")
            (src / "2.lin").write_text(_lin_for(2), encoding="utf-8")
            hands = collect([str(src)])
            self.assertEqual(len(hands), 2)
            files = [f for _, _, fs in hands for f in fs]
            self.assertEqual(len(files), 3)
            self.assertEqual(archive(files, dst), 3)
            self.assertEqual(len(list(dst.glob("*.lin"))), 3)
            self.assertEqual(list(src.glob("*.lin")), [])

    def test_main_writes_docx_and_archives(self):
        with tempfile.TemporaryDirectory() as tmp:
            src, out = Path(tmp, "new"), Path(tmp, "r", "a.docx")
            src.mkdir()
            (src / "1.lin").write_text(_lin_for(3), encoding="utf-8")
            code = main([str(src), "--docx", str(out), "--archive", str(Path(tmp, "done")), "-o", str(Path(tmp, "t.txt"))])
            self.assertEqual(code, 0)
            self.assertTrue(out.exists())
            self.assertEqual(list(src.glob("*.lin")), [])
            self.assertEqual(len(list(Path(tmp, "done").glob("*.lin"))), 1)

    def _docx_text(self, lang):
        lin = _lin_for(4)
        a = analyze_board(parse_lin(lin), lin)
        with tempfile.TemporaryDirectory() as tmp:
            path = build_report([a], [None], str(Path(tmp, "x.docx")), "steffl54",
                                datetime.date(2026, 1, 1), lang=lang)
            doc = Document(path)
            text = "\n".join(p.text for p in doc.paragraphs)
            for t in doc.tables:
                for row in t.rows:
                    text += "\n" + " ".join(c.text for c in row.cells)
        return text

    def test_docx_english_by_default_and_anonymous(self):
        text = self._docx_text("en")
        for secret in ("secretW", "secretN", "secretE"):
            self.assertNotIn(secret, text)
        self.assertIn("steffl54", text)
        for word in ("Auction", "Contract:", "1 January 2026", "Hand analysis"):
            self.assertIn(word, text)
        for ro in ("Licitația", "Analiză", "puncte de onori"):
            self.assertNotIn(ro, text)

    def test_docx_romanian_still_available(self):
        text = self._docx_text("ro")
        self.assertIn("Licitația", text)
        self.assertIn("Analiză done", text)
        self.assertNotIn("Auction", text)

    def test_missing_file_is_skipped_with_a_message(self):
        with tempfile.TemporaryDirectory() as tmp:
            code = main([str(Path(tmp, "nope.lin")), "-o", str(Path(tmp, "t.txt"))])
        self.assertEqual(code, 1)

    def test_result_text_both_languages(self):
        lin = _lin_for(5)
        a = analyze_board(parse_lin(lin), lin)
        a.declarer_tricks = 8
        self.assertEqual(result_text(a), "down 1 (8 tricks)")
        self.assertIn("a căzut 1", result_text(a, "ro"))
        a.declarer_tricks = 9
        self.assertEqual(result_text(a), "made exactly (9 tricks)")
        a.declarer_tricks = 10
        self.assertEqual(result_text(a), "made 10 tricks (+1)")


if __name__ == "__main__":
    unittest.main()
