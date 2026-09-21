"""Word report (English by default, Romanian available): summary page, then one
page per board with the hand diagram, auction, result, double-dummy table and
observations.

Only the chosen username's seat is named; the others are compass directions.
"""
from __future__ import annotations

import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .analysis import SUIT_SYMBOL, BoardAnalysis, side_of, summarize
from .lin_tools import SUITS, auction_grid

RED = RGBColor(0xC0, 0x00, 0x00)
BLACK = RGBColor(0, 0, 0)

GAME_LEVEL = {"NT": 3, "H": 4, "S": 4, "C": 5, "D": 5}

MONTHS_EN = ["January", "February", "March", "April", "May", "June", "July",
             "August", "September", "October", "November", "December"]

LANG = {
    "en": {
        "dir": {"N": "NORTH", "E": "EAST", "S": "SOUTH", "W": "WEST"},
        "side": {"NS": "N-S", "EW": "E-W"},
        "vul": {"None": "None", "NS": "N-S", "EW": "E-W", "Both": "Both"},
        "role": {"declarer": "declarer", "dummy": "dummy", "defender": "defender"},
        "title": "Hand analysis",
        "subtitle": "{date} · player: {user} · {n} boards",
        "headers": ["Board", "Contract", "Result", "Your score", "vs DD (declarer)", "Role"],
        "all_pass": "all pass",
        "res_unknown": "result unknown (the LIN has no card play or claim)",
        "res_exact": "made exactly ({t})",
        "res_over": "made {t} (+{d})",
        "res_down": "down {d} ({t})",
        "short_exact": "exact ({t})",
        "short_over": "{t} (+{d})",
        "short_down": "-{d} ({t})",
        "trick": "trick", "tricks": "tricks",
        "vs_equal": "equal",
        "vs_diff": "{n} {unit} {rel}",
        "above": "above", "below": "below",
        "notes": {
            "slam_light": "Slam with only {hcp} combined HCP (a slam usually needs about 33). Rough rule of thumb.",
            "game_light": "Game with only {hcp} combined HCP (game usually needs about 25). Rough rule of thumb.",
            "part_heavy": "Part-score with {hcp} combined HCP (game values). Rough rule of thumb.",
        },
        "missed_grand": "Double dummy gives {n} tricks: a grand slam was available.",
        "missed_slam": "Double dummy gives {n} tricks: a slam was available.",
        "missed_game": "Double dummy gives {n} tricks in this strain: a game was available.",
        "not_makeable": "The contract could not be made at double dummy (down {n}).",
        "par_other": "The par belonged to {par_side} ({par}, {score}), but the contract was played by {decl_side}.",
        "auction": "Auction",
        "no_auction": "(no auction)",
        "dd_title": "Double dummy (tricks each declarer can take)",
        "par": "Par",
        "observations": "Observations",
        "center": ["Board {n}", "Dealer: {d}", "Vulnerable: {v}"],
        "passed_out": "Contract: passed out.",
        "contract_line": "Contract: {c} played by {decl} — {res}",
        "score_line": "Score: {side} {n}",
        "hcp_line": "HCP: {parts}  (declarer + dummy: {side})",
        "vs_line": "Compared with double dummy: {x}",
        "made_line": "Contracts made as declarer: {m}/{p}",
        "total_line": "Total raw score for your side: {s:+d} over {n} boards",
        "avg_line": "Average HCP held: {a:.1f}",
        "doc_title": "Hand analysis - {date}",
        "date": lambda d: f"{d.day} {MONTHS_EN[d.month - 1]} {d.year}",
    },
    "ro": {
        "dir": {"N": "NORD", "E": "EST", "S": "SUD", "W": "VEST"},
        "side": {"NS": "N-S", "EW": "E-V"},
        "vul": {"None": "nimeni", "NS": "N-S", "EW": "E-V", "Both": "toți"},
        "role": {"declarer": "declarant", "dummy": "mort", "defender": "apărător"},
        "title": "Analiză done",
        "subtitle": "{date} · jucător: {user} · {n} board-uri",
        "headers": ["Board", "Contract", "Rezultat", "Scor pentru tine", "Față de DD (declarant)", "Rol"],
        "all_pass": "pas general",
        "res_unknown": "rezultat necunoscut (LIN-ul nu conține jocul sau claim-ul)",
        "res_exact": "a făcut exact ({t})",
        "res_over": "a făcut {t} (+{d})",
        "res_down": "a căzut {d} ({t})",
        "short_exact": "exact ({t})",
        "short_over": "{t} (+{d})",
        "short_down": "-{d} ({t})",
        "trick": "levată", "tricks": "levate",
        "vs_equal": "egal cu double dummy",
        "vs_diff": "cu {n} {unit} {rel} double dummy",
        "above": "peste", "below": "sub",
        "notes": {
            "slam_light": "Slem cu doar {hcp} puncte de onori împreună (un slem cere de obicei ~33). Regulă orientativă.",
            "game_light": "Joc cu doar {hcp} puncte de onori împreună (jocul cere de obicei ~25). Regulă orientativă.",
            "part_heavy": "Parțial cu {hcp} puncte de onori împreună (valori de joc). Regulă orientativă.",
        },
        "missed_grand": "Double dummy dă {n} levate: se putea licita mare slem.",
        "missed_slam": "Double dummy dă {n} levate: se putea licita slem.",
        "missed_game": "Double dummy dă {n} levate în acest fit/strain: se putea licita joc.",
        "not_makeable": "Contractul nu se putea ține la double dummy (cădere {n}).",
        "par_other": "Parul donei era al perechii {par_side} ({par}, {score}), dar contractul a fost jucat de {decl_side}.",
        "auction": "Licitația",
        "no_auction": "(fără licitație)",
        "dd_title": "Double dummy (levate pe care le poate face fiecare declarant)",
        "par": "Parul donei",
        "observations": "Observații",
        "center": ["Board {n}", "Donator: {d}", "Vulnerabil: {v}"],
        "passed_out": "Contract: toți au spus pas.",
        "contract_line": "Contract: {c} jucat de {decl} — {res}",
        "score_line": "Scor: {side} {n}",
        "hcp_line": "Puncte de onori: {parts}  (declarant + mort: {side})",
        "vs_line": "Față de double dummy: {x}",
        "made_line": "Contracte făcute ca declarant: {m}/{p}",
        "total_line": "Scor brut total pentru perechea ta: {s:+d} pe {n} board-uri",
        "avg_line": "Puncte de onori medii în mână: {a:.1f}",
        "doc_title": "Analiză done - {date}",
        "date": lambda d: f"{d:%d.%m.%Y}",
    },
}


def _L(lang: str) -> dict:
    return LANG[lang]


def _seat(board, pos: str, username: str, lang: str = "en", title: bool = False) -> str:
    name = _L(lang)["dir"][pos]
    name = name.title() if title else name
    return name + (f" ({username})" if pos == board.steffl_pos else "")


def _tricks(n: int, lang: str) -> str:
    L = _L(lang)
    return f"{n} {L['trick'] if n == 1 else L['tricks']}"


def _result_parts(a: BoardAnalysis):
    """(kind, tricks, diff) with kind in exact/over/down, or None if unknown."""
    c = a.board.contract
    if c is None or a.declarer_tricks is None:
        return None
    diff = a.declarer_tricks - (c.level + 6)
    kind = "exact" if diff == 0 else "over" if diff > 0 else "down"
    return kind, a.declarer_tricks, abs(diff)


def result_text(a: BoardAnalysis, lang: str = "en") -> str:
    L = _L(lang)
    if a.board.contract is None:
        return L["all_pass"]
    parts = _result_parts(a)
    if parts is None:
        return L["res_unknown"]
    kind, t, d = parts
    return L["res_" + kind].format(t=_tricks(t, lang), d=d)


def short_result(a: BoardAnalysis, lang: str = "en") -> str:
    L = _L(lang)
    parts = _result_parts(a)
    if parts is None:
        return "—"
    kind, t, d = parts
    return L["short_" + kind].format(t=_tricks(t, lang), d=d)


def vs_dd_text(a: BoardAnalysis, dd: dict, lang: str = "en") -> str | None:
    L = _L(lang)
    c = a.board.contract
    if c is None or a.declarer_tricks is None:
        return None
    diff = a.declarer_tricks - dd["table"][c.declarer][c.strain]
    if diff == 0:
        return L["vs_equal"]
    unit = L["trick"] if abs(diff) == 1 else L["tricks"]
    return L["vs_diff"].format(n=abs(diff), unit=unit, rel=L["above"] if diff > 0 else L["below"])


def par_text(dd: dict, lang: str = "en") -> str:
    par = " ".join(dd["par_contracts"])
    if lang == "ro":  # West is V in Romanian
        par = re.sub(r"(?<=[♣♦♥♠T])W(?=[=+\-\d])", "V", par)
    return par


def observations(a: BoardAnalysis, dd: dict | None, lang: str = "en") -> list[str]:
    L = _L(lang)
    out = []
    c = a.board.contract
    if a.note_kind:
        out.append(L["notes"][a.note_kind].format(hcp=a.side_hcp))
    if dd and c is not None:
        dd_tricks = dd["table"][c.declarer][c.strain]
        need = c.level + 6
        if dd_tricks >= 13 and c.level < 7:
            out.append(L["missed_grand"].format(n=dd_tricks))
        elif dd_tricks >= 12 and c.level < 6:
            out.append(L["missed_slam"].format(n=dd_tricks))
        elif c.level < GAME_LEVEL[c.strain] and dd_tricks >= GAME_LEVEL[c.strain] + 6:
            out.append(L["missed_game"].format(n=dd_tricks))
        if dd_tricks < need:
            out.append(L["not_makeable"].format(n=need - dd_tricks))
        score = dd["par_score"]
        if score != 0 and (side_of(c.declarer) == "NS") != (score > 0):
            out.append(L["par_other"].format(
                par_side=L["side"]["NS" if score > 0 else "EW"], par=par_text(dd, lang),
                score=abs(score), decl_side=L["side"][side_of(c.declarer)]))
    return out


def _shade(cell, hex_color: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(tcPr.makeelement(qn("w:shd"), {qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): hex_color}))


def _hand_cell(cell, board, pos: str, username: str, lang: str) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(_seat(board, pos, username, lang))
    r.bold = True
    for suit in SUITS:
        cards = board.hands[pos][suit]
        line = cell.add_paragraph()
        line.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = line.add_run(f"{SUIT_SYMBOL[suit]} {' '.join(cards) if cards else '—'}")
        run.font.color.rgb = RED if suit in ("H", "D") else BLACK


def _diagram(doc, board, username: str, lang: str) -> None:
    L = _L(lang)
    t = doc.add_table(rows=3, cols=3)
    for pos, (r, c) in {"N": (0, 1), "W": (1, 0), "E": (1, 2), "S": (2, 1)}.items():
        _hand_cell(t.cell(r, c), board, pos, username, lang)
    mid = t.cell(1, 1)
    mid.text = ""
    vul = L["vul"].get(board.vulnerability, board.vulnerability)
    lines = [L["center"][0].format(n=board.number or "?"),
             L["center"][1].format(d=L["dir"][board.dealer].title()),
             L["center"][2].format(v=vul)]
    for i, line in enumerate(lines):
        para = mid.paragraphs[0] if i == 0 else mid.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(line)
        run.italic = True
        run.font.size = Pt(10)


def _auction(doc, board, username: str, lang: str) -> None:
    L = _L(lang)
    doc.add_paragraph().add_run(L["auction"]).bold = True
    rows = auction_grid(board)
    if not rows:
        doc.add_paragraph(L["no_auction"])
        return
    t = doc.add_table(rows=1 + len(rows), cols=4)
    t.style = "Table Grid"
    for i, pos in enumerate(["W", "N", "E", "S"]):
        cell = t.rows[0].cells[i]
        cell.text = _seat(board, pos, username, lang)
        cell.paragraphs[0].runs[0].bold = True
        _shade(cell, "E8EEF7")
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            t.rows[r].cells[c].text = val


def _dd_table(doc, dd: dict, lang: str) -> None:
    L = _L(lang)
    doc.add_paragraph().add_run(L["dd_title"]).bold = True
    cols = ["C", "D", "H", "S", "NT"]
    t = doc.add_table(rows=5, cols=6)
    t.style = "Table Grid"
    for i, x in enumerate(cols, start=1):
        cell = t.cell(0, i)
        cell.text = SUIT_SYMBOL[x]
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if x in ("H", "D"):
            cell.paragraphs[0].runs[0].font.color.rgb = RED
        _shade(cell, "E8EEF7")
    for r, p in enumerate("NESW", start=1):
        t.cell(r, 0).text = L["dir"][p].title()
        for i, x in enumerate(cols, start=1):
            t.cell(r, i).text = str(dd["table"][p][x])
            t.cell(r, i).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    score = dd["par_score"]
    who = L["side"]["NS" if score >= 0 else "EW"]
    doc.add_paragraph(f"{L['par']}: {par_text(dd, lang)}  ({who} {abs(score)})")


def _summary_page(doc, analyses, dds, username: str, date_text: str, lang: str) -> None:
    L = _L(lang)
    doc.add_heading(L["title"], level=0)
    doc.add_paragraph(L["subtitle"].format(date=date_text, user=username, n=len(analyses)))
    t = doc.add_table(rows=1 + len(analyses), cols=6)
    t.style = "Table Grid"
    for i, h in enumerate(L["headers"]):
        cell = t.rows[0].cells[i]
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
        _shade(cell, "E8EEF7")
    for r, (a, dd) in enumerate(zip(analyses, dds), start=1):
        b, c = a.board, a.board.contract
        contract = L["all_pass"] if c is None else f"{c.level}{c.strain_symbol}{'X' * c.doubled} {L['dir'][c.declarer].title()[0]}"
        score = a.steffl_side_score
        row = [
            b.number or "?", contract,
            short_result(a, lang) if c else "—",
            "—" if score is None else f"{score:+d}",
            (vs_dd_text(a, dd, lang) or "—") if dd else "—",
            L["role"].get(a.role, "—"),
        ]
        for i, v in enumerate(row):
            t.rows[r].cells[i].text = v
    s = summarize(analyses)
    doc.add_paragraph()
    if s.declarer_played:
        doc.add_paragraph(L["made_line"].format(m=s.declarer_made, p=s.declarer_played))
    if s.scored:
        doc.add_paragraph(L["total_line"].format(s=s.total_score, n=s.scored))
    if s.avg_hcp is not None:
        doc.add_paragraph(L["avg_line"].format(a=s.avg_hcp))


def build_report(analyses: list[BoardAnalysis], dds: list, path: str, username: str,
                 date, lang: str = "en") -> str:
    """`dds` is a list parallel to `analyses`; an item is a solver dict or None.
    `date` is a datetime.date."""
    L = _L(lang)
    date_text = L["date"](date)
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    doc.styles["Normal"].paragraph_format.space_after = Pt(3)
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Cm(1.8)
        section.left_margin = section.right_margin = Cm(2.2)
    doc.core_properties.title = L["doc_title"].format(date=date_text)
    _summary_page(doc, analyses, dds, username, date_text, lang)

    for a, dd in zip(analyses, dds):
        b, c = a.board, a.board.contract
        doc.add_page_break()
        doc.add_heading(f"Board {b.number or '?'}", level=1)
        _diagram(doc, b, username, lang)
        doc.add_paragraph()
        _auction(doc, b, username, lang)
        doc.add_paragraph()
        if c is None:
            doc.add_paragraph(L["passed_out"])
        else:
            decl = _seat(b, c.declarer, username, lang, title=True)
            doc.add_paragraph(L["contract_line"].format(
                c=f"{c.level}{c.strain_symbol}{'X' * c.doubled}", decl=decl, res=result_text(a, lang)))
            if a.declarer_score is not None:
                signed = a.declarer_score if side_of(c.declarer) == "NS" else -a.declarer_score
                doc.add_paragraph(L["score_line"].format(side=L["side"]["NS" if signed >= 0 else "EW"], n=abs(signed)))
            doc.add_paragraph(L["hcp_line"].format(
                parts=" · ".join(f"{L['dir'][p].title()} {a.hcp[p]}" for p in "NESW"), side=a.side_hcp))
            if dd:
                extra = vs_dd_text(a, dd, lang)
                if extra:
                    doc.add_paragraph(L["vs_line"].format(x=extra))
        if dd:
            _dd_table(doc, dd, lang)
        obs = observations(a, dd, lang)
        if obs:
            doc.add_paragraph().add_run(L["observations"]).bold = True
            for o in obs:
                doc.add_paragraph(o, style="List Bullet")
    for table in doc.tables:  # tight rows so a board fits on one page
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    para.paragraph_format.space_before = Pt(0)
                    para.paragraph_format.space_after = Pt(0)
    doc.save(path)
    return path
